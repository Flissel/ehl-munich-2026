"""Turn the BMW test data in Dataset/ into Rowboat notes (WP1 demo content).

    python demo/dataset_to_notes.py                 # writes demo/example-data/
    python demo/dataset_to_notes.py --out <vault>   # or straight into a vault

Inputs: feedback, customer studies and sales volumes from Dataset/, plus three hand
checked tables in demo/data/ (component -> theme, study item -> theme, option list
facts) and public BMW Group figures (demo/data/annual_report.json).

Output: notes in Rowboat's template (# Title, ## Info, ## Connected to, ## Activity,
## Key facts, ## Open items) under Projects/ and Topics/. Customer feedback is only
used as aggregated counts, never verbatim (spec D8). Re-running is deterministic.
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

log = logging.getLogger("dataset_to_notes")

HERE = Path(__file__).resolve().parent
DATASET = HERE.parent / "Dataset"
DATA = HERE / "data"
OUT = HERE / "example-data"
MARKER = "managed_by: ehl-demo"

PRODUCTS = {
    "F70": {"name": "BMW 1 Series", "body": "Compact hatchback", "sheet": "F70", "feedback": True},
    "G60": {"name": "BMW 5 Series Sedan", "body": "Executive sedan, incl. BMW i5", "sheet": "G60_G68", "feedback": True},
    "G68": {"name": "BMW 5 Series long wheelbase", "body": "Executive sedan, long wheelbase for China", "sheet": "G60_G68", "feedback": False},
    "G70": {"name": "BMW 7 Series", "body": "Luxury sedan, incl. BMW i7", "sheet": "G70", "feedback": True},
}
# Sheet G60_G68 reports both; China is attributed to the long wheelbase G68, all other markets to G60.
SHARED_SPLIT = {"G60": {"EU", "US", "RoW"}, "G68": {"CN"}}
ROADMAP = {"F70": "F70 Roadmap", "G60": "G60 Roadmap", "G68": "G60 Roadmap", "G70": "G70 Roadmap"}
MARKETS = {"EU": "Market Europe", "CN": "Market China", "US": "Market USA", "RoW": "Market Rest of World"}
EUROPE = {"GB", "DE", "FR", "IT", "ES", "SE", "DK", "FI", "NO", "NL", "BE", "AT", "CH", "PL", "IE", "PT"}
CHANNELS = ["Newsletter", "Social Media", "Email"]
NEEDS = ("Wants", "Difficult to Use")  # unmet needs; defects go to quality, not requirements
SKIP_THEMES = {"out of scope", "unclassified"}
SOURCE_HINT = {
    "Source A": "owner reviews, many countries",
    "Source B": "customer care logs, US",
    "Source C": "quality survey, US",
    "Source D": "appeal survey, US",
}


def region(country: str) -> str:
    if country in ("US", "CN"):
        return country
    return "EU" if country in EUROPE else "RoW"


def fmt(n: float) -> str:
    return f"{n:,.0f}"


def pct(part: float, whole: float) -> str:
    return f"{100 * part / whole:.0f} %" if whole else "n/a"


def product_sales(ctx: dict, code: str) -> dict[str, dict[int, int | None]]:
    sales = ctx["sales"][PRODUCTS[code]["sheet"]]
    keep = SHARED_SPLIT.get(code)
    return {m: v for m, v in sales.items() if keep is None or m in keep}


def clean_item(item: str) -> str:
    return item.removeprefix("If applicable: ").strip()


def top(counter: Counter, n: int | None = None) -> list[tuple[str, int]]:
    """Most common first, ties by name, so the output never depends on set order."""
    return sorted(counter.items(), key=lambda x: (-x[1], str(x[0])))[:n]


def link(folder: str, name: str) -> str:
    return f"[[{folder}/{name}]]"


# ---------- loading ----------

def load_csv_map(path: Path) -> dict[str, str]:
    with path.open(encoding="utf-8", newline="") as f:
        return {row[0]: row[1] for row in list(csv.reader(f))[1:]}


def load_feedback(dataset: Path, themes: dict[str, str]) -> list[dict]:
    """One record per (comment ID, component); theme from the lookup table."""
    rows = []
    for code, meta in PRODUCTS.items():
        if not meta["feedback"]:
            continue
        ws = openpyxl.load_workbook(dataset / f"{code}_feedback_hackathon.xlsx", read_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        col = {h: i for i, h in enumerate(header)}
        for r in it:
            if not r[col["ID"]]:
                continue
            component = (r[col["Vfc level2 Name"]] or "no_class_found").strip()
            rows.append({
                "product": code,
                "id": r[col["ID"]],
                "source": r[col["Source"]],
                "country": r[col["Country"]],
                "region": region(r[col["Country"]]),
                "engine": "BEV" if str(r[col["Engine Type"]]).startswith("BEV") else r[col["Engine Type"]],
                "type": r[col["Feedback Type"]] or "Unclassified",
                "component": component,
                "theme": themes.get(component, "unclassified"),
            })
        log.info("%s: %d feedback rows", code, sum(1 for x in rows if x["product"] == code))
    return rows


def load_sales(dataset: Path) -> dict[str, dict[str, dict[int, int | None]]]:
    """sheet -> market code -> year -> volume (None where the cell is not a number)."""
    wb = openpyxl.load_workbook(dataset / "sales_volumes.xlsx", read_only=True)
    sales = {}
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = next(it)
        sheet = {}
        for r in it:
            if r[1]:
                sheet[r[1]] = {int(h[-4:]): (v if isinstance(v, (int, float)) else None)
                               for h, v in zip(header[2:], r[2:])}
        sales[ws.title] = sheet
    return sales


def load_studies(dataset: Path) -> dict:
    """CN/EU: (product, market) -> item -> mean. US: product -> item -> (unhappy, delighted)."""
    wb = openpyxl.load_workbook(dataset / "F70_G60_G68_G70_customer_studies.xlsx", read_only=True, data_only=True)
    rows = list(wb["CN_EU_2025"].iter_rows(values_only=True))
    market_row = next(r for r in rows if r[1] == "China")
    model_row = next(r for r in rows if r[1] and "Series" in str(r[1]))
    cols, market = {}, None
    for i in range(1, len(model_row)):
        market = {"China": "CN", "EU": "EU"}.get(market_row[i], market)
        m = re.search(r"[FG]\d\d", str(model_row[i] or ""))
        if m:
            cols[i] = (m.group(0), market)
    cn_eu: dict[tuple[str, str], dict[str, float]] = defaultdict(dict)
    item = None
    for r in rows:
        if r[0] and r[0] != "Mean":
            item = str(r[0]).strip()
        elif r[0] == "Mean" and item:
            for i, key in cols.items():
                if isinstance(r[i], (int, float)):
                    cn_eu[key][item] = float(r[i])

    rows = list(wb["US_2025"].iter_rows(values_only=True))
    us_cols = {i: re.search(r"[FG]\d\d", str(v)).group(0) for i, v in enumerate(rows[0]) if re.search(r"[FG]\d\d", str(v or ""))}
    us: dict[str, dict[str, tuple[float, float]]] = defaultdict(dict)
    acc: dict[str, dict[str, float]] = {}
    item = None

    def flush():
        if item:
            for p, d in acc.items():
                us[p][item] = (d["bad"], d["top"])

    for r in rows:
        if r[0] and not (r[1] and "%" in str(r[1])):
            flush()
            item, acc = str(r[0]).strip(), {p: {"bad": 0.0, "top": 0.0} for p in us_cols.values()}
        elif r[0] in ("I Hate It", "A Failure", "Unsatisfactory", "Delightful", "I Love It"):
            key = "bad" if r[0] in ("I Hate It", "A Failure", "Unsatisfactory") else "top"
            for i, p in us_cols.items():
                acc[p][key] += float(r[i] or 0)
    flush()
    return {"cn_eu": dict(cn_eu), "us": dict(us)}


# ---------- aggregation ----------

class Stats:
    """Unique comment counts by product, region, theme, component and feedback type."""

    def __init__(self, rows: list[dict]):
        self.rows = rows
        self.comments = {(r["product"], r["id"]): r for r in rows}

    def unique(self, **flt) -> set:
        return {(r["product"], r["id"]) for r in self.rows if all(r[k] == v for k, v in flt.items())}

    def needs(self, **flt) -> set:
        return {(r["product"], r["id"]) for r in self.rows
                if r["type"] in NEEDS and all(r[k] == v for k, v in flt.items())}

    def top_components(self, n: int, **flt) -> list[tuple[str, int]]:
        c = Counter()
        for comp, ids in self._by_component(**flt).items():
            c[comp] = len(ids)
        return sorted(c.items(), key=lambda x: (-x[1], x[0]))[:n]

    def _by_component(self, **flt) -> dict[str, set]:
        out = defaultdict(set)
        for r in self.rows:
            if r["type"] in NEEDS and r["theme"] not in SKIP_THEMES and all(r[k] == v for k, v in flt.items()):
                out[r["component"]].add((r["product"], r["id"]))
        return out

    def theme_needs(self, **flt) -> Counter:
        c = defaultdict(set)
        for r in self.rows:
            if r["type"] in NEEDS and r["theme"] not in SKIP_THEMES and all(r[k] == v for k, v in flt.items()):
                c[r["theme"]].add((r["product"], r["id"]))
        return Counter({t: len(ids) for t, ids in c.items()})

    def counter(self, field: str, **flt) -> Counter:
        return Counter(self.comments[k][field] for k in self.unique(**flt))


def market_contrast(st: Stats, min_needs: int = 60, min_theme: int = 8) -> list[dict]:
    """Which themes owners of the same model raise more often in one market than elsewhere.
    Compared within one model and within Source A (the only multi-country source), so
    differences between models or source types do not pass as market differences."""
    out = []
    for code, meta in PRODUCTS.items():
        if not meta["feedback"]:
            continue
        for m in MARKETS:
            inside = st.theme_needs(product=code, region=m, source="Source A")
            outside = st.theme_needs(product=code, source="Source A") - inside
            n_in, n_out = sum(inside.values()), sum(outside.values())
            if n_in < min_needs or n_out < min_needs:
                continue
            for theme, k in inside.items():
                if k >= min_theme and outside[theme]:
                    idx = (k / n_in) / (outside[theme] / n_out)
                    if idx >= 1.3:
                        out.append({"market": m, "product": code, "theme": theme, "idx": idx, "n": k})
    return sorted(out, key=lambda x: (-x["idx"], x["product"], x["theme"]))


def contrast_line(c: dict, target: str) -> str:
    where = MARKETS[c["market"]].removeprefix("Market ")
    return (f"{link('Topics', target)}: {c['product']} owners in {where} raise it {c['idx']:.1f} times as often "
            f"as {c['product']} owners elsewhere ({fmt(c['n'])} owner reviews, source A)")


def study_contrast(studies: dict, a: tuple[str, str], b: tuple[str, str], label: str) -> str | None:
    x, y = studies["cn_eu"].get(a, {}), studies["cn_eu"].get(b, {})
    diffs = sorted(((x[k] - y[k], k) for k in x.keys() & y.keys()), reverse=True)
    if not diffs:
        return None
    return (f"{label} (study 2025): higher on " + ", ".join(f"{k} {d:+.1f}" for d, k in diffs[:3])
            + "; lower on " + ", ".join(f"{k} {d:+.1f}" for d, k in diffs[-3:][::-1]))


# ---------- note rendering ----------

def note(title: str, info: dict[str, str], connected: list[str], activity: list[str],
         facts: list[str], open_items: list[str], source: str) -> str:
    parts = ["---", f"source: {source}", MARKER, "---", f"# {title}", "", "## Info"]
    parts += [f"**{k}:** {v}" for k, v in info.items()]
    for head, items in (("Connected to", connected), ("Activity", activity),
                        ("Key facts", facts), ("Open items", open_items)):
        parts += ["", f"## {head}"] + ([f"- {x}" for x in items] or ["- none"])
    return "\n".join(parts) + "\n"


def src(ar: dict, key: str) -> str:
    return f"[source]({ar['sources'][key]})"


def study_lines(product: str, studies: dict, item_theme: dict[str, str], theme: str | None = None) -> list[str]:
    lines = []
    for (p, mk), items in sorted(studies["cn_eu"].items()):
        if p != product:
            continue
        sel = {k: v for k, v in items.items() if theme is None or item_theme.get(k) == theme}
        if not sel:
            continue
        ranked = sorted(sel.items(), key=lambda x: (-x[1], x[0]))
        label = "China" if mk == "CN" else "Europe"
        if theme is None:
            lines.append(f"Customer study {label} 2025 (mean, 0 to 10): best {', '.join(f'{clean_item(k)} {v:.1f}' for k, v in ranked[:3])}; "
                         f"weakest {', '.join(f'{clean_item(k)} {v:.1f}' for k, v in ranked[-3:][::-1])}")
        else:
            lines.append(f"{product} study {label} 2025: " + ", ".join(f"{clean_item(k)} {v:.1f}" for k, v in ranked))
    us = studies["us"].get(product, {})
    sel = {k: v for k, v in us.items() if theme is None or item_theme.get(k) == theme}
    if sel:
        worst = sorted(sel.items(), key=lambda x: (-x[1][0], x[0]))[:5 if theme is None else 4]
        lines.append(f"{product} study USA 2025, share unsatisfied or worse: " + ", ".join(f"{clean_item(k)} {v[0] * 100:.0f} %" for k, v in worst))
    return lines


def product_note(code: str, ctx: dict) -> str:
    meta, st, ar, ol = PRODUCTS[code], ctx["stats"], ctx["ar"], ctx["options"][code]
    sales = product_sales(ctx, code)
    shared = " (sheet G60_G68, see assumption)" if code in SHARED_SPLIT else ""
    info = {"Model": f"{meta['name']} ({code})", "Segment": meta["body"],
            "Roadmap": link("Projects", ROADMAP[code]), "Price list": ol["price_list"]}
    if ol["models"]:
        prices = [p for group in ol["models"].values() for p in group.values()]
        info["Drivetrains"] = ", ".join(ol["models"])
        info["Base price Germany"] = f"{fmt(min(prices))} to {fmt(max(prices))} EUR"

    connected = [f"{link('Topics', MARKETS[m])}: {fmt(v[2025])} units 2025{shared}"
                 for m, v in sales.items() if v.get(2025)]
    if code in ("G60", "G68"):
        connected.append(f"{link('Projects', 'G68' if code == 'G60' else 'G60')}: sibling (standard and long wheelbase)")
    if meta["feedback"]:
        for theme, n in top(st.theme_needs(product=code), 6):
            connected.append(f"{link('Topics', theme)}: {fmt(n)} comments with wants or usability issues")
    connected += [link("Topics", "Trend Neue Klasse"), link("Topics", "Market Research")]
    if code != "F70":
        connected.append(link("Topics", "Trend Electrification"))

    facts = []
    for group, models in ol["models"].items():
        facts.append(f"{group}: " + ", ".join(f"{m} {fmt(p)} EUR" for m, p in models.items()))
    facts += ol["tech"]
    if ol["packages"]:
        facts.append("Packages: " + "; ".join(f"{p['name']} ({p['code']}, {fmt(p['price'])} EUR): {p['contents']}" for p in ol["packages"]))
    facts.append(f"Sales{shared}: " + "; ".join(
        f"{MARKETS[m].removeprefix('Market ')} " + " / ".join(fmt(v[y]) if v[y] is not None else "no data" for y in (2024, 2025, 2030))
        for m, v in sales.items()) + " (2024 / 2025 / 2030 plan)")
    if code in SHARED_SPLIT:
        facts.append("Assumption: sheet G60_G68 reports both models; China is attributed to the long wheelbase G68 (built for China), all other markets to G60")
    if meta["feedback"]:
        n = st.unique(product=code)
        types, eng, srcs = (st.counter(f, product=code) for f in ("type", "engine", "source"))
        countries = top(st.counter("country", product=code), 4)
        facts.append(f"Feedback: {fmt(len(n))} unique comments; " + ", ".join(f"{k} {pct(v, len(n))}" for k, v in top(types)))
        facts.append("Feedback by country: " + ", ".join(f"{k} {fmt(v)}" for k, v in countries)
                     + "; by drivetrain: " + ", ".join(f"{k} {fmt(v)}" for k, v in sorted(eng.items())))
        facts.append("Feedback sources (meaning inferred, to be confirmed): " + ", ".join(
            f"{k} {fmt(v)} ({SOURCE_HINT.get(k, 'unknown')})" for k, v in sorted(srcs.items())))
    facts += study_lines(code, ctx["studies"], ctx["item_theme"])
    facts += [f"{t} {src(ar, s)}" for t, s in ar["models"].get(code, [])]

    open_items = []
    if meta["feedback"]:
        areas = [x for x in st.top_components(12, product=code) if ctx["themes"][x[0]] != "Value and overall impression"]
        for comp, n in areas[:6]:
            open_items.append(f"Candidate requirement area: {comp} ({fmt(n)} comments, {ctx['themes'][comp]}); check against the option list")
    else:
        open_items.append("No customer comments in Dataset/ for G68; evidence comes from the China study only")
    if code == "F70":
        open_items.append("Sales sheet F70 China contains a broken formula; treated as no data")
    activity = ["2025: customer studies China, Europe and USA"]
    edition = re.search(r"(\d\d)\.(\d\d)\.(\d{4})", ol["price_list"])
    if edition:
        activity.insert(0, f"{edition.group(3)}-{edition.group(2)}-{edition.group(1)}: price list edition for July 2026")
    return note(code, info, connected, activity, facts, open_items, "dataset")


def roadmap_note(code: str, ctx: dict) -> str:
    sheet = PRODUCTS[code]["sheet"]
    sales = ctx["sales"][sheet]
    members = [c for c, m in PRODUCTS.items() if m["sheet"] == sheet]
    tot = {y: sum(v[y] or 0 for v in sales.values()) for y in (2024, 2025, 2030)}
    info = {"Products": ", ".join(link("Projects", c) for c in members),
            "Horizon": "2024 to 2030 (sales volume plan from Dataset/)",
            "Volume 2030": f"{fmt(tot[2030])} units ({(tot[2030] / tot[2025] - 1) * 100:+.0f} % vs 2025)"}
    connected = [link("Projects", c) for c in members] + [link("Topics", "Trend Neue Klasse")]
    facts, growth = [], []
    for m, v in sales.items():
        if v[2025] and v[2030] is not None:
            g = (v[2030] / v[2025] - 1) * 100
            growth.append((g, m))
            facts.append(f"{MARKETS[m].removeprefix('Market ')}: {fmt(v[2025])} units 2025, {fmt(v[2030])} in 2030 ({g:+.0f} %), share 2030 {pct(v[2030], tot[2030])}")
            connected.append(f"{link('Topics', MARKETS[m])}: {pct(v[2030], tot[2030])} of the 2030 volume")
        else:
            facts.append(f"{MARKETS[m].removeprefix('Market ')}: no volume in the plan")
    growth.sort(reverse=True)
    activity = [f"{y}: {fmt(tot[y])} units" + (" (plan)" if y == 2030 else "") for y in (2024, 2025, 2030)]
    open_items = [f"Fastest growing market to 2030: {MARKETS[growth[0][1]].removeprefix('Market ')} ({growth[0][0]:+.0f} %)",
                  "Requirements are re-prioritised over about two years before the successor is frozen; re-run the converter with every new data wave"]
    return note(ROADMAP[code], info, connected, activity, facts, open_items, "dataset")


def theme_note(theme: str, ctx: dict) -> str:
    st, item_theme = ctx["stats"], ctx["item_theme"]
    comps = sorted(c for c, t in ctx["themes"].items() if t == theme)
    needs_all = st.theme_needs()
    info = {"Type": "Customer theme (from BMW component tags, Dataset/)",
            "Components": f"{len(comps)} BMW component tags, e.g. " + ", ".join(comps[:6])}
    connected, facts = [], []
    for code, meta in PRODUCTS.items():
        if meta["feedback"]:
            n_all = len(st.unique(product=code, theme=theme))
            n_needs = st.theme_needs(product=code).get(theme, 0)
            if n_all:
                connected.append(f"{link('Projects', code)}: {fmt(n_all)} comments, {pct(n_needs, n_all)} wants or usability issues")
                tc = st.counter("type", product=code, theme=theme)
                facts.append(f"{code}: " + ", ".join(f"{k} {fmt(v)}" for k, v in top(tc)))
        facts += study_lines(code, ctx["studies"], item_theme, theme)
    if any(code == "G68" for code, _ in ctx["studies"]["cn_eu"]):
        connected.append(link("Projects", "G68"))
    connected += [contrast_line(c, MARKETS[c["market"]]) for c in ctx["contrast"] if c["theme"] == theme]
    for tname, t in ctx["ar"]["trends"].items():
        if f"Topics/{theme}" in t["links"]:
            connected.append(link("Topics", tname))
    connected.append(link("Topics", "Market Research"))
    open_items = [f"{comp}: {fmt(n)} comments with wants or usability issues"
                  for comp, n in st.top_components(6, theme=theme)]
    info["Unmet needs"] = f"{fmt(needs_all.get(theme, 0))} comments with wants or usability issues across F70, G60 and G70"
    return note(theme, info, connected, [], facts, open_items, "dataset")


def market_note(m: str, ctx: dict) -> str:
    st, ar = ctx["stats"], ctx["ar"]
    countries = st.counter("country", region=m)
    info = {"Region": MARKETS[m].removeprefix("Market "),
            "Feedback countries": ", ".join(f"{k} {fmt(v)}" for k, v in top(countries, 6)) or "none in Dataset/"}
    connected, facts = [], []
    for code, meta in PRODUCTS.items():
        v = product_sales(ctx, code).get(m, {})
        if v.get(2025):
            connected.append(f"{link('Projects', code)}: {fmt(v[2025])} units 2025, {fmt(v[2030] or 0)} in 2030")
    connected += [contrast_line(c, c["theme"]) for c in ctx["contrast"] if c["market"] == m][:4]
    for tname, t in ar["trends"].items():
        if f"Topics/{MARKETS[m]}" in t["links"]:
            connected.append(link("Topics", tname))
    connected.append(link("Topics", "Market Research"))
    facts += [f"{t} {src(ar, s)}" for t, s in ar["markets"].get(m, [])]
    n = st.unique(region=m)
    if n:
        named = top(st.theme_needs(region=m), 4)
        facts.append(f"Feedback: {fmt(len(n))} unique comments; themes with most wants or usability issues: " + ", ".join(f"{t} {fmt(c)}" for t, c in named))
    pairs = {"CN": [(("G70", "CN"), ("G70", "EU"), "G70 in China vs G70 in Europe"),
                    (("G68", "CN"), ("G60", "EU"), "Long wheelbase G68 in China vs G60 in Europe")],
             "EU": [(("G70", "EU"), ("G70", "CN"), "G70 in Europe vs G70 in China")]}
    facts += [line for a, b, label in pairs.get(m, []) if (line := study_contrast(ctx["studies"], a, b, label))]
    open_items = []
    if m == "CN":
        open_items.append("No customer comments from China in Dataset/; only the China study (G68, G70)")
    if m == "US":
        open_items.append("Sources B, C and D are US only, so US evidence is richer than other markets; weight sources before comparing markets")
    if m == "EU":
        open_items.append("European comments are almost only F70 from Great Britain; G60 and G70 evidence for Europe comes from the EU study")
    return note(MARKETS[m], info, connected, [], facts, open_items, "dataset")


def trend_note(name: str, t: dict, ar: dict) -> str:
    return note(name, {"Type": "Trend from public BMW Group sources", "Summary": t["summary"]},
                [link(*x.split("/", 1)) for x in t["links"]] + [link("Topics", "Market Research")],
                [], [f"{x} {src(ar, s)}" for x, s in t["facts"]], [], "bmw-group-public")


def research_note(ctx: dict) -> str:
    ar = ctx["ar"]
    connected = [link("Topics", MARKETS[m]) for m in MARKETS] + [link("Topics", t) for t in ctx["theme_names"]]
    connected += [link("Topics", t) for t in ar["trends"]] + [link("Projects", c) for c in PRODUCTS]
    facts = [f"{t} {src(ar, s)}" for t, s in ar["group"]] + [
        "Dataset/: customer feedback F70, G60, G70 (sources A to D), customer studies 2025 (China, EU, US), sales volumes 2024, 2025 and 2030 plan, German price lists July 2026",
        "Customer themes come from a hand written lookup of BMW component tags (demo/data/themes.csv); study items are mapped the same way (demo/data/study_themes.csv)",
    ]
    open_items = ["Competitor data is not in Dataset/ and has to come from web research",
                  "Ask the organisers what feedback sources A to D are"]
    return note("Market Research", {"Type": "Research hub", "Scope": "Voice of customer, markets and BMW strategy for F70, G60, G68 and G70"},
                connected, [], facts, open_items, "dataset")


def channel_note(name: str) -> str:
    return note(name, {"Type": "Customer communication channel"},
                [link("Topics", "Market Research")] + [link("Projects", c) for c in PRODUCTS], [], [], [], "example")


# ---------- main ----------

def build(dataset: Path) -> dict[str, str]:
    themes = load_csv_map(DATA / "themes.csv")
    ctx = {
        "themes": themes,
        "item_theme": load_csv_map(DATA / "study_themes.csv"),
        "options": json.loads((DATA / "option_lists.json").read_text(encoding="utf-8")),
        "ar": json.loads((DATA / "annual_report.json").read_text(encoding="utf-8")),
        "sales": load_sales(dataset),
        "studies": load_studies(dataset),
        "stats": Stats(load_feedback(dataset, themes)),
    }
    st = ctx["stats"]
    ctx["theme_names"] = sorted(set(themes.values()) - SKIP_THEMES)
    ctx["contrast"] = market_contrast(st)

    notes = {f"Projects/{c}.md": product_note(c, ctx) for c in PRODUCTS}
    notes |= {f"Projects/{ROADMAP[c]}.md": roadmap_note(c, ctx) for c in ("F70", "G60", "G70")}
    notes |= {f"Topics/{t}.md": theme_note(t, ctx) for t in ctx["theme_names"]}
    notes |= {f"Topics/{MARKETS[m]}.md": market_note(m, ctx) for m in MARKETS}
    notes |= {f"Topics/{n}.md": trend_note(n, t, ctx["ar"]) for n, t in ctx["ar"]["trends"].items()}
    notes |= {f"Topics/{c}.md": channel_note(c) for c in CHANNELS}
    notes["Topics/Market Research.md"] = research_note(ctx)
    return notes


def write(notes: dict[str, str], out: Path) -> None:
    """Replace notes this script owns; never touch files without the marker."""
    for old in out.glob("*/*.md"):
        rel = old.relative_to(out).as_posix()
        if rel not in notes and MARKER in old.read_text(encoding="utf-8"):
            old.unlink()
            log.info("removed stale %s", rel)
    for rel, text in notes.items():
        path = out / rel
        if path.exists() and MARKER not in path.read_text(encoding="utf-8"):
            log.warning("skipped %s: not written by this script", rel)
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8", newline="\n")
    log.info("wrote %d notes to %s", len(notes), out)


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dataset", type=Path, default=DATASET)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    write(build(args.dataset), args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
