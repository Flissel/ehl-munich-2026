"""Write Torque's data/quotes.json from the Dataset, so a fresh PC shows the customer comment texts.

Torque ships without comment texts: every quote in dist/bmw.json says "Comment <id>, text only in
the local copy". At start the page loads data/quotes.json, a list of [path, text] pairs, and puts each
text at its path in the BMW data. This script builds that list from Dataset/*_feedback_hackathon.xlsx:
quote id "G70-3713" is the row with ID "G70-3713" in the G70 file. Texts are masked with the research pipeline's
mask_pii (e-mails, phone numbers, names) before they are written.

    python demo/torque/make_quotes.py                      # Dataset/ -> demo/torque/data/quotes.json
    python demo/torque/make_quotes.py --dataset <dir> --out <file>

Needs openpyxl, and pandas for the research masking (both in market/.venv).
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import logging
import sys
from collections.abc import Callable
from pathlib import Path

log = logging.getLogger("torque.quotes")

APP = Path(__file__).resolve().parent
ROOT = APP.parents[1]
PLACEHOLDER_END = "text only in the local copy"
FILE_SUFFIX = "_feedback_hackathon.xlsx"


def placeholders(bmw: object) -> list[tuple[str, list]]:
    """(quote id, path of its text) for every quote whose text is still the placeholder."""
    found: list[tuple[str, list]] = []

    def walk(node: object, path: list) -> None:
        if isinstance(node, dict):
            text = node.get("t")
            if isinstance(text, str) and text.endswith(PLACEHOLDER_END) and "id" in node:
                found.append((str(node["id"]), path + ["t"]))
            for key, value in node.items():
                walk(value, path + [key])
        elif isinstance(node, list):
            for i, value in enumerate(node):
                walk(value, path + [i])

    walk(bmw, [])
    return found


def patches(bmw: object, texts: dict[str, str],
            mask: Callable[[str], str]) -> tuple[list[list], list[str]]:
    """quotes.json entries for every placeholder with a known text, plus the ids without one."""
    out: list[list] = []
    missing: list[str] = []
    for quote_id, path in placeholders(bmw):
        text = texts.get(quote_id)
        if text:
            out.append([path, mask(text)])
        elif quote_id not in missing:
            missing.append(quote_id)
    return out, missing


def load_texts(dataset: Path) -> dict[str, str]:
    """'<series>-<comment ID>' -> comment text, from every *_feedback_hackathon.xlsx in dataset."""
    from openpyxl import load_workbook

    texts: dict[str, str] = {}
    for path in sorted(dataset.glob(f"*{FILE_SUFFIX}")):
        series = path.name[: -len(FILE_SUFFIX)]
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        id_col, text_col = header.index("ID"), header.index("Customer Feedback")
        for row in rows:
            cid, text = row[id_col], row[text_col]
            if cid is None or not text:
                continue
            key = str(cid).strip()
            if not key.startswith(f"{series}-"):  # the Dataset writes "G70-0001"; accept bare numbers too
                key = f"{series}-{key}"
            texts.setdefault(key, str(text).strip())
    return texts


def research_mask() -> Callable[[str], str]:
    """mask_pii from research/scripts/prepare_data.py, the same masking the pipeline applies."""
    spec = importlib.util.spec_from_file_location(
        "prepare_data", ROOT / "research" / "scripts" / "prepare_data.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("research/scripts/prepare_data.py not found")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return lambda text: module.mask_pii(text)[0]


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--dataset", type=Path, default=ROOT / "Dataset")
    ap.add_argument("--bmw", type=Path, default=APP / "dist" / "bmw.json")
    ap.add_argument("--out", type=Path, default=APP / "data" / "quotes.json")
    args = ap.parse_args()

    bmw = json.loads(args.bmw.read_text(encoding="utf-8"))
    entries, missing = patches(bmw, load_texts(args.dataset), research_mask())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(entries, ensure_ascii=False), encoding="utf-8")
    log.info("Wrote %d comment texts to %s", len(entries), args.out)
    if missing:
        log.warning("%d quote ids have no text in %s (e.g. %s)", len(missing), args.dataset,
                    ", ".join(missing[:5]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
