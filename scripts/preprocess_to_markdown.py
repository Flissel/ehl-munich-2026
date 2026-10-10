#!/usr/bin/env python3
"""Preprocess BMW hackathon Excel datasets into Rowboat-friendly Markdown.

This covers the three tabular dataset styles:
- feedback/verbatims
- customer studies
- sales volumes

Option-list PDFs are intentionally left out for a separate parser.
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - only hit on missing setup
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with `pip install openpyxl`."
    ) from exc


DEFAULT_DATASET_DIR = Path("dataset")
DEFAULT_OUTPUT_DIR = Path("outputs/markdown")

FEEDBACK_COLUMNS = {
    "id": "ID",
    "source": "Source",
    "brand": "Brand",
    "sales_description": "Sales Description",
    "country": "Country",
    "vehicle_model": "Derivate (e-code)",
    "engine_type": "Engine Type",
    "feedback_type": "Feedback Type",
    "comment": "Customer Feedback",
    "category_l1": "Vfc level2 3 Name",
    "category_l2": "Vfc level2 Name",
    "category_l3": "Vfc level3 Name",
}


def clean_text(value: Any) -> str:
    """Return a compact string safe for Markdown."""
    if value is None:
        return ""
    text = str(value).replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_markdown_cell(value: Any) -> str:
    text = clean_text(value)
    text = text.replace("|", "\\|")
    text = text.replace("\n", "<br>")
    return text


def format_study_cell(value: Any, row: list[Any], index: int) -> str:
    """Format customer-study percentage rows in human-readable form."""
    text = clean_text(value)
    measure = clean_text(row[1]) if len(row) > 1 else ""
    if index >= 2 and "% of column" in measure.lower():
        try:
            number = float(value)
        except (TypeError, ValueError):
            return text
        return f"{number:.1%}"
    return text


def yaml_value(value: Any) -> str:
    text = clean_text(value)
    text = text.replace('"', '\\"')
    return f'"{text}"'


def slugify(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_") or "unknown"


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def workbook_values(path: Path) -> openpyxl.Workbook:
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def row_is_empty(row: Iterable[Any]) -> bool:
    return all(clean_text(cell) == "" for cell in row)


def sheet_rows(sheet: Any) -> list[list[Any]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and row_is_empty(rows[-1]):
        rows.pop()
    return rows


def frontmatter(**fields: Any) -> str:
    lines = ["---"]
    for key, value in fields.items():
        lines.append(f"{key}: {yaml_value(value)}")
    lines.append("---")
    return "\n".join(lines)


def write_text(path: Path, text: str) -> None:
    ensure_dir(path.parent)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def detect_signal_type(feedback_type: str) -> str:
    normalized = feedback_type.lower().strip()
    if normalized in {"likes", "like"}:
        return "praise"
    if normalized in {"wants", "want"}:
        return "unmet_need"
    if normalized in {"difficult to use", "defect"}:
        return "complaint"
    return "unknown"


def feedback_record_from_row(headers: dict[str, int], row: list[Any], source_file: str) -> dict[str, str]:
    def get(column: str) -> str:
        idx = headers.get(column)
        if idx is None or idx >= len(row):
            return ""
        return clean_text(row[idx])

    record = {
        key: get(column_name) for key, column_name in FEEDBACK_COLUMNS.items()
    }
    record["source_file"] = source_file
    record["source_type"] = "customer_feedback"
    record["signal_type"] = detect_signal_type(record["feedback_type"])
    return record


def combine_feedback_records(records: list[dict[str, str]]) -> list[dict[str, Any]]:
    """Merge repeated comment rows while preserving all tags as evidence."""
    combined: dict[tuple[str, str], dict[str, Any]] = {}

    for record in records:
        key = (record["id"], record["comment"])
        if key not in combined:
            combined[key] = {
                "id": record["id"],
                "source_file": record["source_file"],
                "source_type": record["source_type"],
                "source": record["source"],
                "brand": record["brand"],
                "sales_description": record["sales_description"],
                "country": record["country"],
                "vehicle_model": record["vehicle_model"],
                "engine_type": record["engine_type"],
                "comment": record["comment"],
                "feedback_types": set(),
                "signal_types": set(),
                "classifications": [],
            }

        category_parts = [
            record["category_l1"],
            record["category_l2"],
            record["category_l3"],
        ]
        category = " / ".join(part for part in category_parts if part)
        combined[key]["feedback_types"].add(record["feedback_type"] or "unknown_feedback_type")
        combined[key]["signal_types"].add(record["signal_type"])
        combined[key]["classifications"].append(
            {
                "feedback_type": record["feedback_type"] or "unknown_feedback_type",
                "signal_type": record["signal_type"],
                "category": category or "uncategorized",
            }
        )

    result = []
    for item in combined.values():
        item["feedback_types"] = sorted(item["feedback_types"])
        item["signal_types"] = sorted(item["signal_types"])
        seen = set()
        classifications = []
        for classification in item["classifications"]:
            dedupe_key = (
                classification["feedback_type"],
                classification["signal_type"],
                classification["category"],
            )
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            classifications.append(classification)
        item["classifications"] = classifications
        result.append(item)
    return result


def preprocess_feedback(dataset_dir: Path, output_dir: Path) -> list[Path]:
    output_paths: list[Path] = []
    feedback_files = sorted(dataset_dir.glob("*_feedback_hackathon.xlsx"))

    for path in feedback_files:
        wb = workbook_values(path)
        sheet = wb["Feedback_Explorer"] if "Feedback_Explorer" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = sheet_rows(sheet)
        if not rows:
            continue

        header_row = [clean_text(cell) for cell in rows[0]]
        headers = {name: idx for idx, name in enumerate(header_row) if name}

        records = []
        seen = set()
        for row in rows[1:]:
            record = feedback_record_from_row(headers, row, path.name)
            if not record["id"] and not record["comment"]:
                continue
            dedupe_key = (
                record["id"],
                record["comment"],
                record["category_l1"],
                record["category_l2"],
                record["category_l3"],
            )
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            records.append(record)

        vehicle = clean_text(records[0]["vehicle_model"]) if records else path.stem.split("_")[0]
        combined_records = combine_feedback_records(records)
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for record in combined_records:
            market = record["country"] or "unknown_market"
            grouped[market].append(record)

        lines = [
            frontmatter(
                source_file=path.name,
                source_type="customer_feedback",
                vehicle_model=vehicle,
                source_row_count=len(records),
                unique_comment_count=len(combined_records),
                created_from="scripts/preprocess_to_markdown.py",
            ),
            "",
            f"# {vehicle} Customer Feedback",
            "",
            f"Source file: `{path.name}`",
            f"Source rows: {len(records)}",
            f"Unique comments: {len(combined_records)}",
            "",
        ]

        for market in sorted(grouped):
            lines.extend([f"## Market: {market}", ""])
            for record in grouped[market]:
                lines.extend(
                    [
                        f"### {record['id'] or 'unknown_id'}",
                        "",
                        f"- Source: {record['source'] or 'unknown'}",
                        f"- Vehicle: {record['vehicle_model'] or vehicle}",
                        f"- Market: {record['country'] or 'unknown'}",
                        f"- Engine type: {record['engine_type'] or 'unknown'}",
                        f"- Feedback types: {', '.join(record['feedback_types'])}",
                        f"- Signal types: {', '.join(record['signal_types'])}",
                        "",
                        "Classifications:",
                        "",
                    ]
                )
                for classification in record["classifications"]:
                    lines.append(
                        "- "
                        f"{classification['feedback_type']} "
                        f"({classification['signal_type']}): "
                        f"{classification['category']}"
                    )
                lines.extend(
                    [
                        "",
                        "Customer comment:",
                        "",
                        f"> {record['comment'] or 'No comment text available.'}",
                        "",
                    ]
                )

        out_path = output_dir / "feedback" / f"{slugify(vehicle)}_feedback.md"
        write_text(out_path, "\n".join(lines))
        output_paths.append(out_path)

    return output_paths


def row_to_markdown_table(row: list[Any], width: int, formatter: Any = None) -> str:
    cells = []
    for idx in range(width):
        value = row[idx] if idx < len(row) else ""
        if formatter:
            value = formatter(value, row, idx)
        cells.append(clean_markdown_cell(value))
    return "| " + " | ".join(cells) + " |"


def trim_empty_columns(rows: list[list[Any]]) -> list[list[Any]]:
    if not rows:
        return rows
    max_width = max(len(row) for row in rows)
    non_empty_cols = []
    for idx in range(max_width):
        if any(idx < len(row) and clean_text(row[idx]) for row in rows):
            non_empty_cols.append(idx)
    return [[row[idx] if idx < len(row) else "" for idx in non_empty_cols] for row in rows]


def preprocess_customer_studies(dataset_dir: Path, output_dir: Path) -> list[Path]:
    output_paths: list[Path] = []
    study_files = sorted(dataset_dir.glob("*customer_studies.xlsx"))

    for path in study_files:
        wb = workbook_values(path)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = trim_empty_columns(sheet_rows(sheet))
            rows = [row for row in rows if not row_is_empty(row)]
            if not rows:
                continue

            width = max(len(row) for row in rows)
            lines = [
                frontmatter(
                    source_file=path.name,
                    source_type="customer_study",
                    sheet_name=sheet_name,
                    row_count=len(rows),
                    created_from="scripts/preprocess_to_markdown.py",
                ),
                "",
                f"# Customer Study - {sheet_name}",
                "",
                f"Source file: `{path.name}`",
                f"Sheet: `{sheet_name}`",
                "",
                "| " + " | ".join(f"Column {idx + 1}" for idx in range(width)) + " |",
                "| " + " | ".join("---" for _ in range(width)) + " |",
            ]

            for row in rows:
                lines.append(row_to_markdown_table(row, width, formatter=format_study_cell))

            out_path = output_dir / "customer_studies" / f"customer_study_{slugify(sheet_name)}.md"
            write_text(out_path, "\n".join(lines))
            output_paths.append(out_path)

    return output_paths


def preprocess_sales(dataset_dir: Path, output_dir: Path) -> list[Path]:
    output_paths: list[Path] = []
    sales_files = sorted(dataset_dir.glob("sales_volumes.xlsx"))

    for path in sales_files:
        wb = workbook_values(path)
        lines = [
            frontmatter(
                source_file=path.name,
                source_type="sales_volumes",
                created_from="scripts/preprocess_to_markdown.py",
            ),
            "",
            "# Sales Volumes",
            "",
            f"Source file: `{path.name}`",
            "",
        ]

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = [row for row in sheet_rows(sheet) if not row_is_empty(row)]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            lines.extend(
                [
                    f"## Vehicle: {sheet_name}",
                    "",
                    row_to_markdown_table(rows[0], width),
                    "| " + " | ".join("---" for _ in range(width)) + " |",
                ]
            )
            for row in rows[1:]:
                lines.append(row_to_markdown_table(row, width))
            lines.append("")

        out_path = output_dir / "sales" / "sales_volumes.md"
        write_text(out_path, "\n".join(lines))
        output_paths.append(out_path)

    return output_paths


def run(dataset_dir: Path, output_dir: Path) -> list[Path]:
    ensure_dir(output_dir)
    written = []
    written.extend(preprocess_feedback(dataset_dir, output_dir))
    written.extend(preprocess_customer_studies(dataset_dir, output_dir))
    written.extend(preprocess_sales(dataset_dir, output_dir))
    return written


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert feedback, customer study, and sales Excel files to Markdown."
    )
    parser.add_argument(
        "--dataset-dir",
        type=Path,
        default=DEFAULT_DATASET_DIR,
        help="Folder containing the BMW hackathon dataset files.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Folder where Markdown files should be written.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    written = run(args.dataset_dir, args.output_dir)
    print(f"Wrote {len(written)} Markdown files:")
    for path in written:
        print(f"- {path}")


if __name__ == "__main__":
    main()
